from pathlib import Path

import pytest
from openpyxl import Workbook

from src.pipeline.build_foods import build_food_tables
from src.pipeline.load_workbook import load_workbook

HEADER = ["food", "Caloric Value", "Water", "Protein", "Fat", "Carbohydrates", "Sodium", "Nutrition Density"]


def _fake_workbook(path: Path) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "group_a"
    first.append(HEADER)
    # 150 g serving: 80 + 5 + 1 + 64 grams
    first.append([" ibirayi ", 285, 80.0, 5.0, 1.0, 64.0, 0.2, 5.5])
    first.append(["beans", 120, 60.0, 8.0, 0.5, 21.5, 0.05, 7.0])
    second = workbook.create_sheet("group_b")
    second.append(HEADER)
    second.append(["banana", 90, 75.0, 1.0, 0.3, 23.7, 0.0, 4.0])
    workbook.save(path)


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    path = tmp_path / "fake.xlsx"
    _fake_workbook(path)
    return path


def _outputs(tmp_path: Path) -> dict[str, Path]:
    return {
        "interim_dir": tmp_path / "interim",
        "validation_report": tmp_path / "validation.md",
    }


def test_loader_reads_all_sheets_and_preserves_traceability(workbook: Path, tmp_path: Path) -> None:
    result = load_workbook(workbook, **_outputs(tmp_path))

    assert set(result) == {"group_a", "group_b"}
    assert result["group_a"].loc[0, "food"] == "ibirayi"
    assert result["group_a"].loc[0, "source_sheet"] == "group_a"
    assert result["group_a"].loc[0, "source_row"] == 2
    assert result["group_a"].loc[0, "raw_food"] == " ibirayi "
    assert (tmp_path / "interim" / "group_a.csv").exists()
    assert (tmp_path / "validation.md").exists()


def test_loader_can_skip_a_sheet(workbook: Path, tmp_path: Path) -> None:
    result = load_workbook(workbook, ignored_sheets=("group_b",), **_outputs(tmp_path))
    assert set(result) == {"group_a"}


def test_builder_separates_derived_score_and_writes_sources(workbook: Path, tmp_path: Path) -> None:
    foods, nutrients = build_food_tables(workbook, output_dir=tmp_path / "processed", **_outputs(tmp_path))

    assert len(foods) == 3
    assert "nutrition_density_score" in foods.columns
    assert "nutrition_density" not in nutrients["nutrient_name"].unique()
    assert nutrients["amount_per_100g"].notna().all()
    assert nutrients["source"].notna().all()
    for name in ("foods.csv", "food_nutrients.csv", "cleaning_report.md"):
        assert (tmp_path / "processed" / name).exists()


def test_builder_converts_serving_values_to_per_100g(workbook: Path, tmp_path: Path) -> None:
    foods, nutrients = build_food_tables(workbook, output_dir=tmp_path / "processed", **_outputs(tmp_path))

    ibirayi = foods.loc[foods["english_name"] == "ibirayi"].iloc[0]
    assert ibirayi["serving_g_estimated"] == pytest.approx(150.0)
    values = nutrients[nutrients["food_id"] == ibirayi["food_id"]].set_index("nutrient_name")
    assert values.loc["caloric_value", "amount_per_serving"] == pytest.approx(285)
    assert values.loc["caloric_value", "amount_per_100g"] == pytest.approx(190)
    # sodium is stored in grams in the workbook and published in mg
    assert values.loc["sodium", "unit"] == "mg"
    assert values.loc["sodium", "amount_per_serving"] == pytest.approx(200)
    assert values.loc["sodium", "amount_per_100g"] == pytest.approx(133.3333, rel=1e-3)


def test_builder_does_not_write_into_the_real_data_folder(workbook: Path, tmp_path: Path) -> None:
    real_interim = Path("data/processed/interim")
    before = set(real_interim.glob("*")) if real_interim.exists() else set()
    build_food_tables(workbook, output_dir=tmp_path / "processed", **_outputs(tmp_path))
    after = set(real_interim.glob("*")) if real_interim.exists() else set()
    assert after == before
